"""Pure report builders: dict in, file bytes out. No DB/ORM knowledge here —
routers/export.py gathers the data, these functions just render it."""

import io

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

# fpdf2's core fonts are Latin-1 only, so Cyrillic needs this embedded
# (bundled via the `fonts-dejavu-core` apt package in the Dockerfile).
_DEJAVU_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

STATUS_LABELS = {"planned": "В плане", "confirmed": "Подтверждён", "problem": "Просрочен"}


# --- Excel ------------------------------------------------------------------


def _autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(60, max(12, length + 2))


def build_employee_xlsx(data: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Сводка"
    rows = [
        ("Сотрудник", data["employee"]["full_name"]),
        ("Логин", data["employee"]["login"]),
        ("Направление", data["employee"]["direction"]),
        ("Подразделение", data["employee"]["department"]),
        ("Отчёт сформирован", data["generated_at"].isoformat()),
        (),
        ("Навыков в плане", data["stats"]["total"]),
        ("Подтверждено", data["stats"]["confirmed"]),
        ("Просрочено", data["stats"]["problem"]),
        ("Выполнение плана, %", data["stats"]["completion_pct"]),
    ]
    for row in rows:
        ws.append(row)
    for r in range(1, len(rows) + 1):
        ws.cell(row=r, column=1).font = bold
    _autosize_columns(ws)

    ws_plan = wb.create_sheet("План обучения")
    ws_plan.append(["Навык", "Плановая дата", "Статус", "Дата подтверждения"])
    for cell in ws_plan[1]:
        cell.font = bold
    for item in data["plan_items"]:
        ws_plan.append(
            [
                item["skill"],
                item["planned_date"].isoformat(),
                STATUS_LABELS.get(item["status"], item["status"]),
                item["confirmed_date"].isoformat() if item["confirmed_date"] else "",
            ]
        )
    _autosize_columns(ws_plan)

    ws_meet = wb.create_sheet("Встречи")
    ws_meet.append(["Дата", "Провёл(а)", "Подтверждено навыков", "Отмечено проблем"])
    for cell in ws_meet[1]:
        cell.font = bold
    for m in data["meetings"]:
        ws_meet.append([m["date"].isoformat(), m["conducted_by"], m["confirmed_skills"], m["problems"]])
    _autosize_columns(ws_meet)

    ws_prob = wb.create_sheet("Проблемы")
    ws_prob.append(["Дата встречи", "Навык", "Комментарий", "Статус"])
    for cell in ws_prob[1]:
        cell.font = bold
    for p in data["problems"]:
        ws_prob.append([p["date"].isoformat(), p["skill"] or "Общее", p["comment"], "Решена" if p["resolved"] else "Открыта"])
    _autosize_columns(ws_prob)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_department_xlsx(data: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Сводка"
    rows = [
        ("Подразделение", data["department"]["path"]),
        ("Руководитель", data["department"]["manager"] or "—"),
        ("Отчёт сформирован", data["generated_at"].isoformat()),
        (),
        ("Сотрудников", data["stats"]["people"]),
        ("Выполнение плана, %", data["stats"]["completion_pct"]),
        ("Открытых проблем", data["stats"]["open_problems"]),
    ]
    for row in rows:
        ws.append(row)
    for r in range(1, len(rows) + 1):
        ws.cell(row=r, column=1).font = bold
    _autosize_columns(ws)

    ws_emp = wb.create_sheet("Сотрудники")
    ws_emp.append(["ФИО", "Направление", "Подтверждено / Всего", "Выполнение, %", "Открытых проблем"])
    for cell in ws_emp[1]:
        cell.font = bold
    for e in data["employees"]:
        ws_emp.append([e["full_name"], e["direction"], f"{e['confirmed']} / {e['total']}", e["completion_pct"], e["open_problems"]])
    _autosize_columns(ws_emp)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------


def _new_pdf() -> FPDF:
    pdf = FPDF()
    pdf.add_font("DejaVu", "", _DEJAVU_REGULAR)
    pdf.add_font("DejaVu", "B", _DEJAVU_BOLD)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    return pdf


def build_employee_pdf(data: dict) -> bytes:
    pdf = _new_pdf()

    pdf.set_font("DejaVu", "B", 18)
    pdf.cell(0, 10, "Отчёт по сотруднику", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("DejaVu", "", 11)
    pdf.cell(0, 7, f"Сотрудник: {data['employee']['full_name']} ({data['employee']['login']})", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Направление: {data['employee']['direction']}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Подразделение: {data['employee']['department']}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Сформирован: {data['generated_at'].isoformat()}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("DejaVu", "B", 13)
    pdf.cell(
        0,
        8,
        f"Выполнение плана: {data['stats']['completion_pct']}% ({data['stats']['confirmed']} из {data['stats']['total']})",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    if data["stats"]["problem"]:
        pdf.set_text_color(200, 30, 30)
        pdf.cell(0, 8, f"Просрочено: {data['stats']['problem']}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    pdf.set_font("DejaVu", "B", 13)
    pdf.cell(0, 8, "План обучения", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 10)
    if not data["plan_items"]:
        pdf.cell(0, 6, "План обучения пуст", new_x="LMARGIN", new_y="NEXT")
    for item in data["plan_items"]:
        line = f"- {item['skill']} — план {item['planned_date'].isoformat()}, {STATUS_LABELS.get(item['status'], item['status'])}"
        if item["confirmed_date"]:
            line += f" (подтверждён {item['confirmed_date'].isoformat()})"
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("DejaVu", "B", 13)
    pdf.cell(0, 8, "Встречи", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 10)
    if not data["meetings"]:
        pdf.cell(0, 6, "Встреч ещё не было", new_x="LMARGIN", new_y="NEXT")
    for m in data["meetings"]:
        pdf.multi_cell(
            0,
            6,
            f"- {m['date'].isoformat()} · провёл(а) {m['conducted_by']} · "
            f"подтверждено навыков: {m['confirmed_skills']}, отмечено проблем: {m['problems']}",
            new_x="LMARGIN",
            new_y="NEXT",
        )
    pdf.ln(4)

    if data["problems"]:
        pdf.set_font("DejaVu", "B", 13)
        pdf.cell(0, 8, "Проблемы", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 10)
        for p in data["problems"]:
            status = "решена" if p["resolved"] else "ОТКРЫТА"
            pdf.multi_cell(
                0, 6, f"- [{status}] {p['skill'] or 'Общее'}: {p['comment']} (встреча {p['date'].isoformat()})",
                new_x="LMARGIN", new_y="NEXT",
            )

    return bytes(pdf.output())


def build_department_pdf(data: dict) -> bytes:
    pdf = _new_pdf()

    pdf.set_font("DejaVu", "B", 18)
    pdf.cell(0, 10, "Отчёт по подразделению", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("DejaVu", "", 11)
    pdf.cell(0, 7, f"Подразделение: {data['department']['path']}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Руководитель: {data['department']['manager'] or '—'}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Сформирован: {data['generated_at'].isoformat()}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("DejaVu", "B", 13)
    pdf.cell(
        0,
        8,
        f"Сотрудников: {data['stats']['people']} · выполнение плана: {data['stats']['completion_pct']}%",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    if data["stats"]["open_problems"]:
        pdf.set_text_color(200, 30, 30)
        pdf.cell(0, 8, f"Открытых проблем: {data['stats']['open_problems']}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    pdf.set_font("DejaVu", "B", 13)
    pdf.cell(0, 8, "Сотрудники", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 10)
    for e in data["employees"]:
        line = f"- {e['full_name']} ({e['direction']}) — {e['confirmed']}/{e['total']} ({e['completion_pct']}%)"
        if e["open_problems"]:
            line += f", открытых проблем: {e['open_problems']}"
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())
